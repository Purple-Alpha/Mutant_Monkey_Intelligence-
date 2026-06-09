"""Generate Mutant Monkey branded pitch assets: investor/partner deck (.pptx)
and a synthetic sample-case snapshot (.pdf). All entities fictional.

Run with the throwaway venv:
    /tmp/mm_assets_venv/bin/python "6. Internal_Strategy/Pitch/build_pitch_assets.py"
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

HERE = Path(__file__).resolve().parent
LOGO = Path("/home/socialarchitect/.cursor/projects/home-socialarchitect/assets/mutant_monkey_logo.png")

NAVY = RGBColor(0x10, 0x1F, 0x3A)
TEAL = RGBColor(0x1C, 0x7A, 0x82)
GREEN = RGBColor(0x6F, 0xD8, 0x4A)
GREY = RGBColor(0x44, 0x4A, 0x55)

# reportlab colours
R_NAVY = colors.HexColor("#101F3A")
R_TEAL = colors.HexColor("#1C7A82")
R_GREEN = colors.HexColor("#6FD84A")
R_LIGHT = colors.HexColor("#EEF2F6")


# --------------------------------------------------------------------------
# PowerPoint deck
# --------------------------------------------------------------------------
def add_logo(slide, prs, small=True):
    w = Inches(1.6 if small else 3.2)
    # logo aspect ~1.5:1
    if small:
        slide.shapes.add_picture(str(LOGO), Inches(11.4), Inches(0.2), width=w)
    else:
        slide.shapes.add_picture(str(LOGO), Inches(5.0), Inches(1.1), width=w)


def add_title(slide, text, top=0.35, size=30):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(10.5), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return box


def add_bullets(slide, items, top=1.6, left=0.7, width=12.0, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for level, text, *opt in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size - level * 2)
        r.font.color.rgb = GREY if level else NAVY
        if opt and opt[0] == "accent":
            r.font.color.rgb = TEAL
            r.font.bold = True
        p.space_after = Pt(6)


def build_deck(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — Title
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(LOGO), Inches(4.4), Inches(0.7), width=Inches(4.5))
    t = s.shapes.add_textbox(Inches(0.6), Inches(4.5), Inches(12.1), Inches(2.0))
    tf = t.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "Where we are — and the data we need to move forward"
    r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = NAVY
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = "Governance-first threat detection  |  Confidential  |  June 2026"
    r2.font.size = Pt(15); r2.font.color.rgb = TEAL

    # 2 — The problem
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "The problem we attack")
    add_bullets(s, [
        (0, "Business Email Compromise (BEC) and vendor-payment fraud are the most expensive attack class for small and mid-size businesses."),
        (0, "The damage isn't malware — it's a believable email that reroutes a real payment."),
        (1, "A vendor 'updates' their bank details. Finance complies. The money is gone."),
        (1, "Generic spam filters pass it — there's no link, no attachment, no payload."),
        (0, "SMBs lack a SOC. They need detection that is automatic, explainable, and defensible.", "accent"),
    ])

    # 3 — What we're building
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "What we're building")
    add_bullets(s, [
        (0, "A governed swarm of specialist detection agents — not one black-box model."),
        (0, "Five layers, each with a strict job:"),
        (1, "L2 Detection  ·  L3 Verification  ·  L4 Evidence  ·  L5 Challenge  ·  Decision"),
        (0, "Every agent is facts-only and lift-only: it contributes signals, it never acts on its own."),
        (0, "Every agent is operator-signed, audit-gated, and deterministic.", "accent"),
    ])

    # 4 — Where we are now
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "Where we are right now")
    add_bullets(s, [
        (0, "13 governed agents live — up from 8 in the latest build session."),
        (1, "First governed Layer 3 (Verification) and Layer 4 (Evidence) both shipped."),
        (1, "Detection agents now safely ride stateful vendor-baseline memory."),
        (0, "1,410 automated tests passing at Evidence Stage 1 (Synthetic)."),
        (0, "Maturity ladder: Stage 1 Synthetic  →  Stage 2 Supervised  →  Stage 3 Production.", "accent"),
        (1, "We are at Stage 1, fully governed. Stage 2 is gated on real-world data — see slide 8."),
    ])

    # 5 — How it works
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "How it works (the discipline)")
    add_bullets(s, [
        (0, "Proven wrapper pattern, repeated agent after agent:"),
        (1, "Agent → AgentContribution → SwarmCommander → Decision Evidence Record → Layer 5"),
        (0, "No detector logic is ever hand-edited. No agent self-registers. No autonomy."),
        (0, "Each readiness boundary passes an independent audit gate before commit."),
        (0, "Each capability is locked by an operator signature — nothing ships unsigned.", "accent"),
    ])

    # 6 — Why different / moat
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "Why we're worth watching")
    add_bullets(s, [
        (0, "Governance-first is the moat, not a checkbox."),
        (1, "Deterministic + fully auditable: every verdict traces to signed, gated facts."),
        (1, "Explainable by construction — we show why, not just a score."),
        (1, "Safe by design: facts-only agents can't take a harmful action."),
        (0, "This is what regulated buyers and insurers will demand. We built it in from day one.", "accent"),
    ])

    # 7 — Live case (SMB)
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "Live case: 'SMB' — Sterling Mill & Building Co.")
    add_bullets(s, [
        (0, "Synthetic walkthrough (see the one-page snapshot PDF):"),
        (1, "Inbound email: vendor requests a bank-detail change on invoice #SMB-4471."),
        (0, "What the swarm surfaced — as facts, not actions:"),
        (1, "#14 Payment Change Detection: financial-state delta vs. learned baseline."),
        (1, "#11 Known-Good Contact: requested account differs from verified contact record."),
        (1, "#48 Verification Outcome: out-of-band confirmation not completed."),
        (0, "Verdict: routed to human verification before any payment. Loss prevented.", "accent"),
    ])

    # 8 — Data we need
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "The data we need to move forward")
    add_bullets(s, [
        (0, "To graduate from Stage 1 (Synthetic) to Stage 2 (Supervised), we need real signal:"),
        (1, "Design-partner mailflow (read-only) — real vendor/finance email patterns."),
        (1, "Labeled incident data — confirmed BEC / invoice-fraud cases, anonymized."),
        (1, "Vendor payment-change events with known-good outcomes for baseline truth."),
        (0, "In return, partners get early detection plus a say in what we build next.", "accent"),
    ])

    # 9 — Watch or partner
    s = prs.slides.add_slide(blank); add_logo(s, prs); add_title(s, "Watch us — or partner with us")
    add_bullets(s, [
        (0, "Watch: the governed-swarm architecture is the defensible way to do AI security."),
        (0, "Partner (design partner): "),
        (1, "You bring anonymized mailflow / incident data."),
        (1, "You get Stage 2 detection on your own traffic and influence over the roadmap."),
        (0, "The ask: a handful of design partners and a data-sharing conversation.", "accent"),
    ])

    # 10 — Close
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(LOGO), Inches(4.8), Inches(1.6), width=Inches(3.8))
    t = s.shapes.add_textbox(Inches(0.6), Inches(4.8), Inches(12.1), Inches(1.6))
    tf = t.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "Mutant Monkey — governed threat detection, built to be trusted."
    r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = "Let's talk about data access and a design partnership."
    r2.font.size = Pt(16); r2.font.color.rgb = TEAL

    prs.save(out_path)
    return out_path


# --------------------------------------------------------------------------
# Sample case PDF
# --------------------------------------------------------------------------
def _header_footer(canvas, doc):
    canvas.saveState()
    # Slim text-only masthead (no cramped top logo).
    canvas.setFillColor(R_NAVY)
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(0.6 * inch, LETTER[1] - 0.72 * inch, "MUTANT MONKEY")
    canvas.setFillColor(R_TEAL)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawRightString(LETTER[0] - 0.6 * inch, LETTER[1] - 0.7 * inch,
                           "THREAT DETECTION SNAPSHOT")
    canvas.setFillColor(colors.HexColor('#444A55'))
    canvas.setFont("Helvetica", 7.5)
    canvas.drawRightString(LETTER[0] - 0.6 * inch, LETTER[1] - 0.87 * inch,
                           "CONFIDENTIAL  ·  SYNTHETIC SAMPLE")
    canvas.setStrokeColor(R_LIGHT)
    canvas.line(0.6 * inch, LETTER[1] - 1.02 * inch, LETTER[0] - 0.6 * inch, LETTER[1] - 1.02 * inch)

    # Branded logo sign-off — fills the open lower third of the page, centered.
    logo_w = 3.4 * inch
    logo_h = logo_w / 1.5
    canvas.drawImage(str(LOGO), (LETTER[0] - logo_w) / 2.0, 1.15 * inch,
                     width=logo_w, height=logo_h,
                     preserveAspectRatio=True, mask='auto')

    # footer
    canvas.setStrokeColor(R_LIGHT)
    canvas.line(0.6 * inch, 0.85 * inch, LETTER[0] - 0.6 * inch, 0.85 * inch)
    canvas.setFont("Helvetica-Oblique", 7.5)
    canvas.setFillColor(colors.HexColor('#444A55'))
    canvas.drawString(0.6 * inch, 0.68 * inch,
                      "Synthetic sample — all entities, names, and figures are fictional. Generated by Mutant Monkey for demonstration only.")
    canvas.drawRightString(LETTER[0] - 0.6 * inch, 0.68 * inch, "mutantmonkey.ai")
    canvas.restoreState()


def build_pdf(out_path):
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=R_NAVY, fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=R_TEAL, fontSize=12, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10, leading=14, textColor=colors.HexColor('#222831'))
    small = ParagraphStyle("small", parent=body, fontSize=8.5, textColor=colors.HexColor('#444A55'))

    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER,
                            topMargin=1.3 * inch, bottomMargin=3.55 * inch,
                            leftMargin=0.6 * inch, rightMargin=0.6 * inch)
    flow = []
    flow.append(Paragraph("Vendor Payment-Change Review", h1))
    flow.append(Paragraph("Client: <b>Sterling Mill &amp; Building Co.</b> (\u201cSMB\u201d) &nbsp;|&nbsp; Case #SMB-4471 &nbsp;|&nbsp; Stage 1 (Synthetic)", small))

    flow.append(Paragraph("Scenario", h2))
    flow.append(Paragraph(
        "An inbound email \u2014 appearing to come from a known vendor, Harbour Steel Supply \u2014 asked SMB\u2019s finance "
        "team to update the bank account on the next payment against invoice #SMB-4471 ($48,250.00). The message had "
        "no link and no attachment, so conventional spam filtering passed it. The governed swarm reviewed it and "
        "contributed the following facts.", body))

    flow.append(Paragraph("Swarm findings (facts only \u2014 no autonomous action)", h2))
    data = [
        ["Agent", "Layer", "Finding (fact)"],
        ["#14 Payment Change Detection", "L2 Detection", "Requested bank details differ from the learned vendor baseline (financial-state delta)."],
        ["#11 Known-Good Contact", "L3 Verification", "New account does not match any verified known-good contact record for this vendor."],
        ["#48 Verification Outcome", "L3 Verification", "No completed out-of-band (two-channel) confirmation found for this change."],
        ["#46 Evidence Package", "L4 Evidence", "Facts bundled into a single auditable case record for human review."],
    ]
    tbl = Table(data, colWidths=[1.9 * inch, 1.1 * inch, 4.3 * inch])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), R_NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, R_LIGHT]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor('#CBD5E0')),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    flow.append(tbl)

    flow.append(Paragraph("Verdict", h2))
    flow.append(Paragraph(
        "<b>Hold &amp; verify.</b> The payment was routed to mandatory human verification via a known-good phone "
        "contact before release. No funds moved on the changed details. The swarm made no payment decision and took "
        "no autonomous action \u2014 it surfaced signed, gated facts and escalated to a person.", body))

    flow.append(Spacer(1, 8))
    flow.append(Paragraph(
        "Why this matters: the same facts that protect SMB also create an audit trail an insurer or regulator can "
        "trust. Every line above traces to an operator-signed, independently gated detector.", small))

    doc.build(flow, onFirstPage=_header_footer, onLaterPages=_header_footer)
    return out_path


if __name__ == "__main__":
    deck = build_deck(HERE / "MutantMonkey_Investor_Partner_Deck.pptx")
    pdf = build_pdf(HERE / "MutantMonkey_Sample_Case_SMB.pdf")
    print("WROTE:", deck)
    print("WROTE:", pdf)
